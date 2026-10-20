"""
Genera il SAL (Stato Avanzamento Lavori) dettagliato agosto–settembre 2026 in Word.

Requisiti: python-docx in un ambiente virtuale
    python3 -m venv .venv && .venv/bin/pip install python-docx
    .venv/bin/python docs/report/genera_sal_ago_set.py
Dati: docs/report/sal_ago_set_dati.py (una riga per attività, da aggiornare lì).
Output: docs/report/SAL-Cruscotto-HR-Agosto-Settembre-2026.docx
"""
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

sys.path.insert(0, str(Path(__file__).parent))
from genera_resoconto_ago_set import BLU, GRIGIO, LOGO, heading, table  # noqa: E402
from sal_ago_set_dati import AREE  # noqa: E402

OUT = Path(__file__).with_name("SAL-Cruscotto-HR-Agosto-Settembre-2026.docx")
DOC_MARKERS = ("MD", "PDF", ".sql", ".sh", "README", "DEPLOY_", "GUIDA", "wiki", "Stato e Roadmap", "version.json")


def main():
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
    sec.left_margin = sec.right_margin = Cm(1.8)
    sec.top_margin = sec.bottom_margin = Cm(1.6)
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)

    fp = sec.footer.paragraphs[0]
    fp.text = "Cruscotto HR · SAL agosto – settembre 2026"
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in fp.runs:
        r.font.size = Pt(8)
        r.font.color.rgb = GRIGIO

    tot = sum(len(a[2]) for a in AREE)
    completate = sum(1 for a in AREE for x in a[2] if x[4].startswith("Completato"))

    # ── Copertina
    if LOGO.exists():
        doc.add_picture(str(LOGO), width=Cm(12))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()
    for txt, size, color, bold in [
        ("Cruscotto HR", 28, BLU, True),
        ("Stato Avanzamento Lavori (SAL)", 18, None, False),
        ("Periodo: 25 agosto – 29 settembre 2026", 13, GRIGIO, False),
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(txt)
        r.font.size = Pt(size)
        r.bold = bold
        if color:
            r.font.color.rgb = color
    doc.add_paragraph()
    table(
        doc,
        ["Voce", "Dettaglio"],
        [
            ("Prodotto", "Cruscotto HR — Sistema di Monitoraggio HR della Pubblica Amministrazione"),
            ("Oggetto", "Elenco dettagliato delle attività svolte nel periodo"),
            ("Attività censite", f"{tot} (di cui {completate} completate)"),
            ("Aree di intervento", str(len(AREE))),
        ],
        [5, 20],
    )
    doc.add_page_break()

    # ── Riepilogo per area
    heading(doc, "1. Riepilogo per area", 1)
    doc.add_paragraph(
        "Le attività sono raggruppate per area e numerate (es. D.05 = quinta attività "
        "dell'area D). Per ogni attività sono indicati data, descrizione, risultato prodotto "
        "e stato."
    )
    righe = []
    for code, titolo, items in AREE:
        date = [x[0] for x in items]
        righe.append((code, titolo, str(len(items)), f"{date[0]} – {date[-1]}" if len(set(date)) > 1 else date[0]))
    righe.append(("", "Totale", str(tot), "25/08 – 29/09"))
    table(doc, ["Area", "Descrizione", "N. attività", "Periodo"], righe, [1.5, 15, 3, 5])
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.add_run("Legenda stato. ").bold = True
    p.add_run(
        "«Completato»: attività conclusa e verificata. «Consegnato – esecuzione a cura del "
        "sistemista»: lavoro pronto e consegnato, da installare sull'infrastruttura. "
        "«Completato – in attesa di dati/fonti»: predisposto, in attesa delle fonti dati."
    )
    doc.add_page_break()

    # ── Dettaglio
    heading(doc, "2. Dettaglio delle attività", 1)
    for code, titolo, items in AREE:
        heading(doc, f"{code}. {titolo}", 2)
        rows = [
            (f"{code}.{i:02d}", d, att, det, out, st)
            for i, (d, att, det, out, st) in enumerate(items, 1)
        ]
        table(doc, ["N.", "Data", "Attività", "Descrizione", "Risultato / Output", "Stato"], rows,
              [1.4, 1.5, 4.6, 9.6, 5.0, 3.5])
        doc.add_paragraph()

    # ── Documenti consegnati
    doc.add_page_break()
    heading(doc, "3. Documenti e script consegnati", 1)
    docs = [
        (f"{code}.{i:02d}", d, out, att)
        for code, _, items in AREE
        for i, (d, att, _det, out, _st) in enumerate(items, 1)
        if any(m in out for m in DOC_MARKERS)
    ]
    table(doc, ["Rif.", "Data", "Documento / file", "Attività"], docs, [1.6, 1.8, 9.5, 12.6])

    # ── Attività con dipendenze esterne
    doc.add_paragraph()
    heading(doc, "4. Attività con passi a carico di terzi", 1)
    ext = [
        (f"{code}.{i:02d}", att, st)
        for code, _, items in AREE
        for i, (_d, att, _det, _out, st) in enumerate(items, 1)
        if not st == "Completato"
    ]
    table(doc, ["Rif.", "Attività", "Stato"], ext, [1.6, 14, 10])

    doc.save(OUT)
    print(f"Creato {OUT} — {tot} attività")


if __name__ == "__main__":
    main()
