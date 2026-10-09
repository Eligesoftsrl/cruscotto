import sys,re
# Requisito: pip install python-docx
# Uso: python3 docs/report/md2docx.py docs/wiki/07-Vista-Sintetica-Indici-Score.md docs/report/Rapporto-Vista-Sintetica-Indici-Score.docx
from docx import Document
from docx.shared import Pt, RGBColor
src,dst=sys.argv[1],sys.argv[2]
lines=open(src,encoding='utf-8').read().split('\n')
doc=Document()
st=doc.styles['Normal']; st.font.name='Calibri'; st.font.size=Pt(10)
def add_runs(p,text):
    text=re.sub(r'\[([^\]]+)\]\([^)]+\)',r'\1',text)
    parts=re.split(r'(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*)',text)
    for part in parts:
        if not part: continue
        if part.startswith('**'): r=p.add_run(part[2:-2]); r.bold=True
        elif part.startswith('`'): r=p.add_run(part[1:-1]); r.font.name='Consolas'; r.font.color.rgb=RGBColor(0x1F,0x4E,0x79)
        elif part.startswith('*') and len(part)>2: r=p.add_run(part[1:-1]); r.italic=True
        else: p.add_run(part.replace('\\|','|'))
i=0
while i<len(lines):
    l=lines[i]
    if l.startswith('```'):
        i+=1; buf=[]
        while i<len(lines) and not lines[i].startswith('```'): buf.append(lines[i]); i+=1
        p=doc.add_paragraph(); r=p.add_run('\n'.join(buf)); r.font.name='Consolas'; r.font.size=Pt(8.5)
    elif l.startswith('|'):
        rows=[]
        while i<len(lines) and lines[i].startswith('|'):
            cells=[c.strip() for c in re.split(r'(?<!\\)\|',lines[i].strip())[1:-1]]
            if not all(re.fullmatch(r'-+',c) for c in cells if c): rows.append(cells)
            i+=1
        i-=1
        n=max(len(r) for r in rows)
        t=doc.add_table(rows=len(rows),cols=n); t.style='Light Grid Accent 1'
        for ri,r in enumerate(rows):
            for ci in range(n):
                cell=t.cell(ri,ci); cell.text=''
                p=cell.paragraphs[0]; add_runs(p,r[ci] if ci<len(r) else '')
                for run in p.runs:
                    run.font.size=Pt(8.5)
                    if ri==0: run.bold=True
        doc.add_paragraph()
    elif l.startswith('#'):
        lvl=len(l)-len(l.lstrip('#')); doc.add_heading(l.lstrip('#').strip(),level=min(lvl,4) if lvl>1 else 0)
    elif l.startswith('> '):
        p=doc.add_paragraph(); add_runs(p,l[2:]); 
        for r in p.runs: r.italic=True
    elif re.match(r'^\s*[-*] ',l):
        p=doc.add_paragraph(style='List Bullet'); add_runs(p,re.sub(r'^\s*[-*] ','',l))
    elif re.match(r'^\s*\d+\. ',l):
        p=doc.add_paragraph(style='List Number'); add_runs(p,re.sub(r'^\s*\d+\. ','',l))
    elif l.strip()=='---' or not l.strip(): pass
    else:
        p=doc.add_paragraph(); add_runs(p,l.strip())
    i+=1
doc.save(dst); print('saved',dst)
