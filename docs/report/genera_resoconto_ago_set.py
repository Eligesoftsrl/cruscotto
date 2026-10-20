"""
Genera il documento Word «Resoconto attività Agosto–Settembre 2026».

Requisiti: python-docx (pip install python-docx), da eseguire in un ambiente virtuale:
    python3 -m venv .venv && .venv/bin/pip install python-docx
    .venv/bin/python docs/report/genera_resoconto_ago_set.py
Output: docs/report/Resoconto-Attivita-Agosto-Settembre-2026.docx
"""
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BLU = RGBColor(0x00, 0x40, 0x80)
GRIGIO = RGBColor(0x55, 0x5F, 0x6B)
OUT = Path(__file__).with_name("Resoconto-Attivita-Agosto-Settembre-2026.docx")
LOGO = Path(__file__).resolve().parents[2] / "src" / "assets" / "header-logos.png"

# ─────────────────────────────────────────────────────────────────────────────
# CONTENUTI
# ─────────────────────────────────────────────────────────────────────────────
SINTESI = [
    "Acquisizione del prototipo e sua trasformazione in un'applicazione strutturata, "
    "manutenibile e verificabile.",
    "Collegamento alla base dati reale: 15 viste di compatibilità costruite e validate sui "
    "dati effettivi, con mappatura completa di grafici, tabelle e chiavi ente.",
    "Conto Annuale interamente su dati reali: 12 schede su 12, con filtri dinamici a cascata "
    "e selettore ente.",
    "Accesso con credenziali aziendali (single sign-on) e profilazione: DFP, "
    "amministratore, responsabile HR di ente, con perimetro dei dati per codice fiscale.",
    "Strato di protezione dei dati tra applicazione e base dati, attivato in staging.",
    "Pannello di amministrazione operativo: attivazione delle schede, statistiche d'uso e "
    "registri di accessi, eventi ed errori.",
    "Meccanismo di aggiornamento sicuro delle versioni e identità grafica (font, favicon).",
    "Documentazione: report tecnici per committente e sistemisti e wiki di progetto con "
    "manuale utente.",
]

NUMERI = [
    ("Interventi registrati nel periodo", "oltre 120"),
    ("Schede Conto Annuale collegate ai dati reali", "12 su 12"),
    ("Viste dati di compatibilità prodotte e validate", "15 (più 13 viste predisposte)"),
    ("Tabelle mappate rispetto ai grafici", "54"),
    ("Componenti riorganizzati sul livello di servizio unico", "46"),
    ("Servizi di accesso ai dati creati", "18"),
    ("File più grande dell'applicazione", "da 1.189 a 340 righe"),
    ("Caricamento iniziale", "da pacchetto unico di ~1,9 MB a pagine a richiesta"),
    ("Errori rilevati dai controlli automatici del codice", "da 7 a 0"),
    ("Documenti tecnici consegnati", "oltre 15 (Markdown e PDF)"),
]

# ── Dal prototipo all'applicazione ──────────────────────────────────────────
PARTENZA = [
    ("Accesso ai dati sparso", "La base dati veniva interrogata direttamente da circa 46 "
     "componenti grafici: ogni modifica a una tabella o a una query costringeva a toccare "
     "decine di file e non era possibile verificare la logica in modo isolato."),
    ("Dati finti", "Molte sezioni mostravano numeri inventati scritti nel codice, non i dati "
     "reali."),
    ("Indirizzi e chiavi nel codice", "Indirizzi del database scritti a mano: l'applicazione "
     "non funzionava fuori dal computer dello sviluppatore e non poteva passare da sviluppo a "
     "produzione senza modifiche al codice."),
    ("Logica mescolata alla grafica", "Interrogazioni e calcoli dentro i componenti di "
     "interfaccia: codice difficile da leggere, correggere e riutilizzare."),
    ("Duplicazione", "Griglie di indicatori, stati di caricamento e stili dei grafici copiati "
     "in ogni sezione: ogni ritocco andava ripetuto a mano."),
    ("File troppo grandi", "Alcuni file facevano «tutto» (fino a 1.189 righe)."),
    ("Nessun controllo di qualità", "Assenza di regole automatiche e test: facile introdurre "
     "regressioni senza accorgersene."),
    ("Caricamento lento", "Un unico pacchetto di circa 1,9 MB da scaricare alla prima "
     "apertura."),
    ("Nessuna gestione degli errori", "Un errore in un grafico poteva lasciare la pagina "
     "bianca; gli errori di caricamento non venivano segnalati."),
    ("Accesso dimostrativo", "Scelta del profilo da un elenco, senza autenticazione reale né "
     "perimetro dei dati."),
]

LIVELLI = [
    ("Componente (presentazione)", "Mostra i dati: card, grafici, tabelle, filtri. Non sa da "
     "dove arrivano i dati né come si calcolano.", "La card «Personale in servizio»"),
    ("Hook (orchestrazione)", "Collega il componente ai dati: legge i filtri attivi, avvia il "
     "caricamento, gestisce attesa, errori e memoria temporanea (cache), evita richieste "
     "ripetute.", "«Dammi i dati della scheda Età per i filtri scelti»"),
    ("Service (accesso ai dati)", "Unico punto che parla con la base dati: prepara i parametri "
     "(anno, ente, comparto…), chiama le funzioni di calcolo, trasforma il risultato nel "
     "formato atteso dall'interfaccia.", "Chiamata alla funzione di calcolo dell'età e "
     "conversione dei campi"),
    ("Configurazione", "Indirizzi e impostazioni letti dall'ambiente (sviluppo, staging, "
     "produzione), mai scritti nel codice.", "Indirizzo della base dati e di Keycloak"),
    ("Strato di protezione e base dati", "Verifica l'identità dell'utente e restituisce solo "
     "i dati del suo perimetro.", "Un responsabile HR vede solo il proprio ente"),
]

PRIMA_DOPO = [
    ("Accesso ai dati", "Dentro ~46 componenti", "Centralizzato in 18 servizi"),
    ("Dati delle sezioni", "In gran parte dimostrativi", "Reali; i dimostrativi residui sono "
     "segnalati"),
    ("Configurazione", "Indirizzi scritti nel codice", "Un'unica configurazione per ambiente"),
    ("Interfaccia", "Elementi duplicati in ogni sezione", "Componenti riutilizzabili e tema "
     "comune"),
    ("Dimensione dei file", "Fino a 1.189 righe", "Il più grande ridotto a 340 righe"),
    ("Caricamento", "Pacchetto unico ~1,9 MB", "Pagine caricate a richiesta"),
    ("Errori", "Pagina bianca, errori silenziosi", "Protezione dagli errori e avvisi all'utente"),
    ("Qualità", "Nessun controllo", "Regole automatiche (0 errori), formattazione su 175 file, "
     "test"),
    ("Tipi di dato", "Generici, non controllati", "Servizi principali tipizzati sulla base "
     "dati"),
    ("Accesso", "Profilo scelto da elenco", "Single sign-on Keycloak, ruoli e perimetro per "
     "ente"),
]

DIFFICOLTA = [
    ("Due anagrafiche enti diverse", "Nella base dati esistevano due elenchi di enti con "
     "identificativi sovrapposti (gli stessi numeri indicavano enti diversi).",
     "Scelta di un'anagrafica unica e, in seguito, del codice fiscale dell'ente come chiave "
     "definitiva."),
    ("Dati in forma «a dizionario»", "La base dati reale memorizza gli indicatori come coppie "
     "elemento/valore, mentre l'applicazione si aspetta tabelle per ente.",
     "Viste di trasformazione che riorganizzano i dati senza modificare l'interfaccia "
     "(264 enti, 78 metriche)."),
    ("Codici diversi tra tabelle", "Le causali nei dati usano codici con prefisso (es. A23, "
     "C01) diversi da quelli delle tabelle di decodifica.", "Tabelle di decodifica dedicate, "
     "verificate sui dati reali."),
    ("Fonti non ancora disponibili", "InPA, Syllabus, competenze e PTFP non erano presenti "
     "negli estratti forniti.", "Viste predisposte vuote e avviso «dati a scopo dimostrativo» "
     "nelle sezioni interessate, in attesa delle fonti."),
    ("Accesso Keycloak bloccato dal browser", "Il controllo silenzioso della sessione in un "
     "riquadro nascosto veniva bloccato dalle politiche di sicurezza del server di identità.",
     "Passaggio al solo flusso con reindirizzamento, più robusto."),
    ("Token senza informazioni sull'ente", "Il token di accesso inizialmente non indicava "
     "l'ente dell'utente.", "Contratto dei claim concordato con chi gestisce Keycloak: ruoli e "
     "codice fiscale dell'ente (enti_cf)."),
    ("Orologi dei server non allineati", "Lo strato di protezione rifiutava token validi "
     "perché l'orologio del server era avanti/indietro di alcuni minuti.", "Tolleranza "
     "configurabile e richiesta di sincronizzazione oraria al sistemista."),
    ("Versioni vecchie nei browser", "Gli utenti continuavano a vedere versioni precedenti "
     "(anche l'icona) a causa della memoria del browser.", "Identificativo di versione, regole "
     "di cache e banner «Nuova versione disponibile»."),
    ("Errori dopo sessioni scadute", "Sessioni vecchie generavano errori di autorizzazione "
     "diffusi.", "Le chiamate amministrative e il tracciamento partono solo dopo un accesso "
     "valido."),
]

SICUREZZA = [
    ("Nessuna credenziale nel codice", "Indirizzi e chiavi letti dalla configurazione; "
     "verificato che nessun segreto finisca nel pacchetto scaricato dal browser."),
    ("Accesso diretto ai dati vietato", "Una regola automatica blocca qualsiasi componente che "
     "provi a interrogare la base dati fuori dai servizi."),
    ("Accesso aziendale sicuro", "Single sign-on con protocollo PKCE, senza segreti "
     "nell'applicazione; segnalata la necessità di rigenerare un segreto esposto e di usare un "
     "client pubblico."),
    ("Ruoli consentiti espliciti", "Solo i ruoli previsti accedono; gli altri vedono la "
     "schermata «Accesso non autorizzato»."),
    ("Perimetro dei dati per ente", "Ogni responsabile HR vede solo i dati del proprio ente "
     "(codice fiscale dal token); senza ente associato nessun dato aggregato."),
    ("Separazione dei privilegi", "La vista su tutti gli enti (DFP) non comporta l'accesso al "
     "pannello di amministrazione."),
    ("Strato di protezione dei dati", "Servizio intermedio che verifica il token e applica le "
     "regole di accesso prima di interrogare la base dati; regole di accesso per riga (RLS) "
     "predisposte per il sistemista."),
    ("Dati personali", "Segnalato come priorità massima il blocco dell'accesso anonimo alle "
     "tabelle con dati di login."),
    ("Registri sicuri", "I registri non contengono dettagli tecnici sensibili; esportazioni "
     "protette dall'inserimento di formule malevole nei fogli di calcolo."),
    ("Risorse solo interne", "Font istituzionale incluso nell'applicazione: nessuna richiesta a "
     "servizi esterni, compatibile con le politiche di sicurezza dei contenuti."),
]

# Ogni area: (titolo, periodo, introduzione, [(attività, dettaglio)])
AREE = [
    (
        "1. Acquisizione e ristrutturazione dell'applicazione",
        "25 – 27 agosto",
        "Il prototipo iniziale è stato acquisito, reso funzionante nell'ambiente di sviluppo e "
        "riorganizzato per poter crescere in modo ordinato.",
        [
            ("Acquisizione del codice", "Importazione del progetto, avvio dell'applicazione e "
             "analisi iniziale della struttura e delle dipendenze."),
            ("Configurazione per ambienti", "Un unico punto di configurazione per gli indirizzi "
             "dei servizi, senza valori scritti nel codice; porte e ricaricamento adattati "
             "all'ambiente locale e a quello gestito."),
            ("Livello di servizio unico", "Creati 18 servizi che centralizzano l'accesso ai dati; "
             "12 componenti di caricamento dati semplificati; 46 componenti di sezione e grafico "
             "non accedono più direttamente alla base dati."),
            ("Componenti riutilizzabili", "Mattoni comuni per indicatori, griglie, card dei "
             "grafici, stati di caricamento/errore/vuoto, tema dei grafici e formattazione dei "
             "numeri: interfaccia uniforme in tutte le schede."),
            ("Tipizzazione dei dati", "I servizi principali del Conto Annuale sono stati "
             "tipizzati con i tipi generati dalla base dati, eliminando i tipi generici."),
            ("Primo collegamento ai dati reali", "Analisi per età, Benchmark, Cessazioni, Genere, "
             "Progressioni, Lavoro agile e flessibile, Tasso di turnover e di sostituzione "
             "collegati alle tabelle reali; i dati ancora dimostrativi sono contrassegnati."),
            ("Pulizia", "Rimozione di codice inutilizzato e di file non pertinenti al prodotto; "
             "nome del pacchetto e istruzioni di avvio (README) aggiornati."),
        ],
    ),
    (
        "2. Qualità, robustezza e prestazioni",
        "26 – 27 agosto",
        "Interventi per rendere l'applicazione più stabile, veloce e controllabile.",
        [
            ("Protezione dagli errori", "Componente che impedisce lo «schermo bianco» quando una "
             "parte della pagina va in errore; gestione centralizzata degli errori di "
             "caricamento con messaggi all'utente."),
            ("Regole di qualità automatiche", "Controlli automatici che impediscono l'accesso "
             "diretto alla base dati fuori dai servizi; errori di analisi statica azzerati."),
            ("Stile del codice uniforme", "Formattazione automatica e regole comuni applicate a "
             "175 file."),
            ("Prestazioni", "Caricamento delle pagine a richiesta (suddivisione del pacchetto) e "
             "impostazioni di cache dei dati più efficienti."),
            ("Test automatici", "Primi test unitari sulle trasformazioni dei dati di età, genere "
             "e cessazioni."),
            ("Correzioni", "Chiavi duplicate nei KPI, grafici radar senza dati, avvisi di "
             "navigazione eliminati."),
        ],
    ),
    (
        "3. Autenticazione e profilazione degli utenti",
        "2 – 29 settembre",
        "Introdotto l'accesso con le credenziali aziendali (Keycloak) e definita la regola che "
        "stabilisce chi vede quali dati.",
        [
            ("Accesso single sign-on", "Login con redirect e protocollo sicuro (PKCE), senza "
             "segreti nell'applicazione; rinnovo automatico della sessione; verificato end-to-end "
             "in locale con l'ambiente reale."),
            ("Solo accesso aziendale", "Con Keycloak attivo la pagina di accesso mostra soltanto "
             "il pulsante «Accedi con SSO»: ruolo ed ente arrivano dal profilo aziendale."),
            ("Mappatura dei ruoli", "Riconoscimento dei ruoli amministrativi (DFP, Formez, "
             "amministratore unico) e del responsabile HR di ente; tabella dei profili allineata "
             "alle combinazioni di ruoli reali."),
            ("Perimetro per ente", "Lettura dell'ente dal token: prima tramite codice ISTAT, poi "
             "in modo definitivo tramite codice fiscale dell'ente (claim «enti_cf»), con "
             "verifica sui dati reali (Comune di Roma, Sapienza)."),
            ("Protezione dalle fughe di dati", "Un responsabile HR senza ente associato non vede "
             "dati aggregati nazionali; un ruolo non riconosciuto viene bloccato."),
            ("Separazione DFP / amministratore", "La vista su tutti gli enti non dà più accesso al "
             "pannello di amministrazione: regola applicata sia nell'applicazione sia nello strato "
             "di protezione dei dati."),
            ("Selettore ente per il DFP", "Ricerca e selezione di una singola amministrazione, "
             "riservata agli utenti DFP."),
            ("Utenti non autorizzati", "Schermata dedicata «Accesso non autorizzato» con "
             "pulsante Esci, invece del pulsante di accesso."),
            ("Contratto dei claim", "Documento operativo per chi gestisce Keycloak con i dati "
             "da inserire nel token (ruoli, codice fiscale ente)."),
        ],
    ),
    (
        "4. Allineamento alla base dati reale",
        "3 – 8 settembre",
        "Studio della base dati reale e costruzione di viste che la rendono compatibile con "
        "l'applicazione, senza modificare l'interfaccia.",
        [
            ("Mappatura grafici ↔ tabelle", "Per ogni grafico e sezione: tabelle utilizzate, "
             "chiave ente e colonne lette (54 tabelle), con scheletro SQL delle viste da "
             "ricreare."),
            ("Studio della base dati reale", "Analisi della struttura a dizionario (64 elementi e "
             "27 indicatori), individuazione del codice fiscale come chiave ente e piano di "
             "allineamento."),
            ("Anagrafica enti", "Vista degli enti costruita e verificata: 23.594 enti."),
            ("Indicatori KPI", "Trasformazione dei dati a dizionario in tabella per ente "
             "(264 enti monitorati, 78 metriche) e viste di rilevazione KPI."),
            ("Viste del Conto Annuale", "Sei viste occupazionali dai Conti Annuali reali: "
             "occupazione (1,2 milioni di righe), assunzioni, cessazioni, età (2,8 milioni), "
             "formazione, lavoro flessibile; collegamento agli enti tramite codice fiscale "
             "(12.274 corrispondenze verificate)."),
            ("Tabelle di decodifica", "Causali di assunzione/cessazione, comparti e contratti, "
             "fasce d'età e qualifiche per esteso (corrispondenza verificata al 100%), passaggi "
             "di qualifica."),
            ("Viste predisposte", "13 viste conformi ma ancora vuote per le fonti non disponibili "
             "(InPA, Syllabus, competenze, PTFP), verificate su una base dati pulita."),
            ("Strumenti per i colleghi", "Script per l'esportazione della base dati e documento "
             "«vista → tabelle sorgenti → query» con tutte le interrogazioni."),
        ],
    ),
    (
        "5. Conto Annuale su dati reali",
        "10 – 29 settembre",
        "Tutte le 12 schede del Conto Annuale sono state collegate alle funzioni di calcolo "
        "della base dati, secondo le mappature condivise.",
        [
            ("Analisi per età", "Prima scheda collegata e verificata a schermo; definito lo "
             "schema replicato poi su tutte le altre."),
            ("Anzianità di servizio", "Quattro indicatori, distribuzione ed evoluzione per "
             "genere."),
            ("Assunti per causale", "Indicatori e grafici su dati reali; rimossa una voce di "
             "menu non più prevista."),
            ("Cessazioni dal servizio", "Indicatori e grafici su dati reali."),
            ("Tasso di turnover e di sostituzione", "Ultime due schede del primo blocco."),
            ("Sei schede del secondo blocco", "Progressioni, Analisi del personale, Lavoro "
             "flessibile, Formazione, Analisi per genere e Lavoro agile, verificate una per una "
             "rispetto alle mappature."),
            ("Rifiniture richieste", "Filtro Genere bloccato dove non applicabile, titoli senza "
             "anno fisso, note e etichette chiarite, grafico non previsto rimosso."),
            ("Filtri dinamici a cascata", "Opzioni dei filtri lette dai dati, con azzeramento "
             "a catena; nuovo ordine Comparto → Macrocategoria → Categoria."),
            ("Documentazione", "Riepilogo tecnico della scheda Età, documentazione tecnica delle "
             "schede e nota di metodo; sei file di mappatura."),
        ],
    ),
    (
        "6. Navigazione e dati dimostrativi",
        "14 – 22 settembre",
        "Interventi sull'esperienza di navigazione e sulla chiarezza dei dati non ancora reali.",
        [
            ("Tasto «Indietro» del browser", "La navigazione tra le sezioni è registrata nella "
             "cronologia: il tasto Indietro riporta alla vista precedente."),
            ("Correzione di una pagina non trovata", "Risolto l'errore sulla scheda «Assunti per "
             "causale»."),
            ("Modalità «N/D»", "Interruttore centrale che mostra «non disponibile» al posto dei "
             "valori dimostrativi; rettifiche ai menu."),
            ("Avviso dati dimostrativi", "Banner «Dati a scopo dimostrativo» su tutte le sezioni "
             "ancora non collegate (InPA, Minerva, Syllabus, Lavoro Pubblico, D1 e D3); nessun "
             "avviso sul Conto Annuale, interamente reale."),
        ],
    ),
    (
        "7. Pannello di amministrazione",
        "17 – 24 settembre",
        "Area riservata agli amministratori per governare le funzionalità e monitorare l'uso.",
        [
            ("Gestione delle schede", "Attivazione e disattivazione delle schede con effetto "
             "immediato per tutti gli utenti, salvata nella base dati."),
            ("Statistiche d'uso", "Grafico delle schede attive/disattive e delle schede più "
             "consultate."),
            ("Registri", "Registrazione reale di accessi, eventi ed errori (senza dati tecnici "
             "sensibili), consultazione e svuotamento."),
            ("Schema dati", "Script SQL delle tabelle di amministrazione aggiornato."),
            ("Guida tecnica", "Documento per replicare l'area di amministrazione su altri "
             "progetti (schema, raccolta dati, sicurezza, checklist)."),
        ],
    ),
    (
        "8. Sicurezza: strato di protezione dei dati",
        "23 – 25 settembre",
        "Servizio intermedio tra applicazione e base dati che verifica l'identità dell'utente "
        "e applica il perimetro dei dati.",
        [
            ("Pacchetto pronto al rilascio", "Servizio di protezione con configurazione, "
             "verifiche di sicurezza e istruzioni per il sistemista."),
            ("Regole sui dati (RLS)", "Bozza delle regole di accesso per ente e adeguamento delle "
             "funzioni di calcolo, con sequenza consigliata di attivazione."),
            ("Aggancio «a interruttore»", "L'applicazione usa il servizio di protezione solo "
             "quando è configurato, senza impatti sugli altri ambienti."),
            ("Messa in esercizio in staging", "Diagnosi e soluzione dei problemi di avvio: "
             "tolleranza per piccoli sfasamenti dell'orologio dei server, correzione SQL "
             "suggerita al sistemista, chiamate amministrative eseguite solo dopo l'accesso."),
        ],
    ),
    (
        "9. Rilascio, versioni e identità grafica",
        "21 – 25 settembre",
        "Accorgimenti perché gli utenti vedano sempre la versione più recente.",
        [
            ("Aggiornamento sicuro delle versioni", "Identificativo univoco a ogni rilascio, "
             "regole di cache per il server web e banner «Nuova versione disponibile — "
             "Ricarica»."),
            ("Icona dell'applicazione", "Stemma come favicon, con aggiornamento forzato nei "
             "browser."),
            ("Font Titillium Web", "Font istituzionale incluso nell'applicazione, senza "
             "richieste a servizi esterni (compatibile con le politiche di sicurezza)."),
        ],
    ),
    (
        "10. Documentazione e wiki di progetto",
        "26 agosto – 25 settembre",
        "Documenti per committente, esperti e sistemisti, e wiki pubblicabile.",
        [
            ("Report di ristrutturazione", "Report delle attività con esempi «prima/dopo», "
             "schermate e mappa dei dati dimostrativi, in versione definitiva e «attività in "
             "corso», senza riferimenti a strumenti esterni."),
            ("Domande tecniche al committente", "Sette aree di domande (enti/comuni, "
             "protezione dei dati, Keycloak, amministrazione, percorsi guidati, dati, rilascio) "
             "con i punti bloccanti evidenziati."),
            ("Documenti per i sistemisti", "Mappature dati, viste SQL e loro sorgenti, contratto "
             "dei claim, regole di accesso, istruzioni per lo strato di protezione."),
            ("Wiki di progetto", "Manuale utente, ruoli e accessi, pannello di gestione, "
             "funzionalità, architettura tecnica e pagina «Stato e attività» con le note di "
             "rilascio; registro formale, menu laterale e collegamenti nel formato della wiki."),
        ],
    ),
]

CRONOLOGIA = [
    ("25/08", "Acquisizione del prototipo, configurazione per ambienti, livello di servizio unico."),
    ("26/08", "Primo collegamento ai dati reali, componenti riutilizzabili, regole di qualità, "
              "prestazioni, test, report di ristrutturazione."),
    ("27/08", "Protezione dagli errori, stile del codice uniforme, documento di domande tecniche."),
    ("31/08", "Revisione dei report per il committente (versione «attività in corso»)."),
    ("02/09", "Accesso con credenziali aziendali (Keycloak) e mappatura dei ruoli."),
    ("03/09", "Lettura dell'ente dal token; mappatura grafici ↔ tabelle e scheletro delle viste."),
    ("07/09", "Studio della base dati reale; viste enti, KPI, Conto Annuale e decodifiche."),
    ("08/09", "Viste predisposte, documento viste/sorgenti/query, verifica dei contenuti."),
    ("10/09", "Contratto dei claim; scheda Analisi per età su dati reali; selettore ente DFP."),
    ("12/09", "Schede Anzianità, Assunti, Cessazioni, Turnover e Sostituzione su dati reali."),
    ("14/09", "Tasto Indietro, modalità N/D, correzioni ai menu, nota di metodo, mappature."),
    ("16/09", "Sei schede del secondo blocco del Conto Annuale su dati reali."),
    ("17/09", "Rifiniture schede; sezione «Schede» del pannello di amministrazione; perimetro "
              "per codice fiscale."),
    ("18/09", "Pannello di amministrazione collegato alla base dati (schede e registri)."),
    ("21–22/09", "Favicon, aggiornamento sicuro delle versioni, ruoli consentiti, avvisi dati "
                 "dimostrativi."),
    ("23/09", "Strato di protezione dei dati, regole di accesso, aggancio a interruttore, wiki."),
    ("24/09", "Wiki riscritta (manuale e architettura), profili, separazione DFP/amministratore, "
              "guida tecnica dell'area amministrazione."),
    ("25/09", "Messa in esercizio in staging, font istituzionale, favicon, note di rilascio."),
    ("29/09", "Schermata per utenti non autorizzati; nuovo ordine dei filtri del Conto Annuale."),
]


# ─────────────────────────────────────────────────────────────────────────────
# IMPAGINAZIONE
# ─────────────────────────────────────────────────────────────────────────────
def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for r in h.runs:
        r.font.color.rgb = BLU
        r.font.name = "Calibri"
    return h


def table(doc, headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = ""
        run = c.paragraphs[0].add_run(h)
        run.bold = True
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        shade(c, "004080")
    for k, row in enumerate(rows):
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(v)
            run.font.size = Pt(10)
            if i == 0:
                run.bold = True
            if k % 2 == 1:
                shade(cells[i], "EEF3F8")
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Cm(w)
    # intestazione ripetuta su ogni pagina
    trPr = t.rows[0]._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)
    return t


def main():
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Cm(2.2))
    sec.top_margin = sec.bottom_margin = Cm(2.0)

    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)

    # Piè di pagina
    fp = sec.footer.paragraphs[0]
    fp.text = "Cruscotto HR · Resoconto attività agosto – settembre 2026"
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in fp.runs:
        r.font.size = Pt(8)
        r.font.color.rgb = GRIGIO

    # Copertina
    if LOGO.exists():
        doc.add_picture(str(LOGO), width=Cm(12))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Cruscotto HR")
    r.bold = True
    r.font.size = Pt(30)
    r.font.color.rgb = BLU
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Resoconto delle attività di sviluppo")
    r.font.size = Pt(18)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Agosto – Settembre 2026")
    r.font.size = Pt(14)
    r.font.color.rgb = GRIGIO
    for _ in range(8):
        doc.add_paragraph()
    table(
        doc,
        ["Voce", "Dettaglio"],
        [
            ("Prodotto", "Cruscotto HR — Sistema di Monitoraggio HR della Pubblica Amministrazione"),
            ("Periodo", "25 agosto – 29 settembre 2026"),
            ("Contenuto", "Elenco dettagliato delle attività svolte, per area e in ordine cronologico"),
        ],
        [4, 12.5],
    )
    doc.add_page_break()

    # 1. Sintesi
    heading(doc, "Sintesi", 1)
    doc.add_paragraph(
        "Nel periodo il Cruscotto è passato da prototipo dimostrativo ad applicazione collegata "
        "ai dati reali, con accesso aziendale, profilazione per ente, area di amministrazione e "
        "strato di protezione dei dati. I risultati principali:"
    )
    for s in SINTESI:
        doc.add_paragraph(s, style="List Bullet")
    doc.add_paragraph()
    heading(doc, "Il periodo in numeri", 2)
    table(doc, ["Indicatore", "Valore"], NUMERI, [11.5, 5])
    doc.add_page_break()

    # 2. Dal prototipo all'applicazione
    heading(doc, "Dal prototipo all'applicazione", 1)
    doc.add_paragraph(
        "Il punto di partenza era un prototipo realizzato rapidamente per mostrare "
        "l'idea: valido per far vedere il funzionamento, ma non adatto a essere mantenuto "
        "ed evoluto nel tempo né a lavorare con dati reali e utenti reali. Il lavoro di "
        "agosto e settembre lo ha trasformato in un'applicazione strutturata, mantenendo "
        "invariato ciò che l'utente vede."
    )
    heading(doc, "La situazione di partenza", 2)
    table(doc, ["Problema", "Perché era un problema"], PARTENZA, [4.5, 12])
    doc.add_paragraph()

    heading(doc, "La nuova architettura a livelli", 2)
    doc.add_paragraph(
        "L'applicazione è stata organizzata in livelli con responsabilità separate. Ogni "
        "livello parla solo con quello successivo: il percorso di un dato è sempre "
        "Componente → Hook → Service → strato di protezione → base dati."
    )
    table(doc, ["Livello", "Cosa fa", "Esempio"], LIVELLI, [4, 8.5, 4])
    doc.add_paragraph()
    doc.add_paragraph(
        "Perché conta: se cambia una tabella o una funzione di calcolo si interviene su un "
        "solo servizio, non su decine di schermate; la logica si può verificare con test "
        "automatici; le nuove schede si costruiscono riutilizzando gli stessi mattoni. Una "
        "regola automatica impedisce di aggirare questa organizzazione."
    )
    heading(doc, "Confronto prima / dopo", 2)
    table(doc, ["Aspetto", "Prototipo", "Applicazione attuale"], PRIMA_DOPO, [3.8, 5.6, 7.1])
    doc.add_paragraph()

    heading(doc, "Difficoltà incontrate e soluzioni", 2)
    table(doc, ["Difficoltà", "Descrizione", "Soluzione adottata"], DIFFICOLTA, [3.8, 6.4, 6.3])
    doc.add_paragraph()

    heading(doc, "Sicurezza", 2)
    doc.add_paragraph(
        "La sicurezza è stata affrontata su più livelli: codice, accesso, dati e "
        "infrastruttura."
    )
    table(doc, ["Intervento", "Cosa garantisce"], SICUREZZA, [4.5, 12])
    doc.add_page_break()

    # 3. Attività per area
    heading(doc, "Attività svolte per area", 1)
    for titolo, periodo, intro, righe in AREE:
        heading(doc, titolo, 2)
        p = doc.add_paragraph()
        r = p.add_run(f"Periodo: {periodo}")
        r.italic = True
        r.font.color.rgb = GRIGIO
        doc.add_paragraph(intro)
        table(doc, ["Attività", "Dettaglio"], righe, [4.5, 12])
        doc.add_paragraph()

    doc.add_page_break()
    # 4. Cronologia
    heading(doc, "Cronologia", 1)
    doc.add_paragraph("Principali attività per giornata di lavoro.")
    table(doc, ["Data", "Attività"], CRONOLOGIA, [2.5, 14])

    doc.save(OUT)
    print(f"Creato {OUT}")


if __name__ == "__main__":
    main()
